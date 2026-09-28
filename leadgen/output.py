"""Write leads to Google Sheets (upsert by Lead ID) and to local CSV/XLSX."""
from __future__ import annotations

import csv
import json
import os
from pathlib import Path

from .common import SHEET_COLUMNS

# Columns the sales team edits by hand — never overwritten on re-runs.
MANUAL_COLUMNS = {"Status", "Notes"}


def write_csv(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=SHEET_COLUMNS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows: list[dict], path: Path, extra_tabs: dict[str, list[dict]] | None = None) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    tabs = {"Leads": rows, **(extra_tabs or {})}
    for i, (title, tab_rows) in enumerate(tabs.items()):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        _fill_sheet(ws, tab_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _fill_sheet(ws, rows: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    ws.append(SHEET_COLUMNS)
    for r in rows:
        ws.append([r.get(c, "") for c in SHEET_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    widths = {"Company": 24, "Evidence Quote": 60, "Evidence URL": 40, "Matching Case Study": 45,
              "Suggested Pitch": 45, "Pain Signals": 35, "Notes": 35}
    for i, col in enumerate(SHEET_COLUMNS, 1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 14)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    colors = {"Hot": "F8D7DA", "Warm": "FFF3CD", "Cold": "E2E3E5",
              "High": "F8D7DA", "Medium": "FFF3CD", "Low": "E2E3E5"}
    for name in ("Priority", "Recent Volume"):
        idx = SHEET_COLUMNS.index(name)
        for row in ws.iter_rows(min_row=2):
            fill = colors.get(row[idx].value)
            if fill:
                row[idx].fill = PatternFill("solid", fgColor=fill)
    status_col = ws.cell(row=1, column=SHEET_COLUMNS.index("Status") + 1).column_letter
    dv = DataValidation(type="list", formula1='"New,Contacted,Replied,Meeting Booked,Won,Lost,Not a Fit"')
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{max(len(rows) + 1, 500)}")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions


def write_google_sheet(rows: list[dict], sheet_id: str, tab: str = "Leads") -> str:
    """Upsert rows into a Google Sheet via a service account.

    Auth: GOOGLE_SERVICE_ACCOUNT_JSON (the key file's JSON content) or GOOGLE_APPLICATION_CREDENTIALS
    (a path). Share the sheet with the service account's email as Editor.
    """
    try:
        import gspread
    except ImportError:
        raise RuntimeError("Google Sheets support isn't installed. Run: pip install -r requirements-sheets.txt "
                           "(or leave 'Update Google Sheet' unticked and use Download .xlsx)") from None

    raw = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON")
    gc = gspread.service_account_from_dict(json.loads(raw)) if raw else gspread.service_account(
        filename=os.getenv("GOOGLE_APPLICATION_CREDENTIALS"))
    sh = gc.open_by_key(sheet_id)
    try:
        ws = sh.worksheet(tab)
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(tab, rows=1000, cols=len(SHEET_COLUMNS))

    existing = ws.get_all_records() if ws.row_count and ws.acell("A1").value else []
    by_id = {str(r.get("Lead ID")): r for r in existing}
    for r in rows:
        old = by_id.get(r["Lead ID"])
        if old:  # keep the sales team's edits and the original discovery date
            for c in MANUAL_COLUMNS | {"Date Found"}:
                if old.get(c):
                    r[c] = old[c]
        by_id[r["Lead ID"]] = r
    merged = sorted(by_id.values(), key=lambda r: int(r.get("Lead Score") or 0), reverse=True)

    ws.clear()
    ws.update([SHEET_COLUMNS] + [[r.get(c, "") for c in SHEET_COLUMNS] for r in merged],
              value_input_option="USER_ENTERED")
    ws.freeze(rows=1, cols=5)
    header = f"A1:{gspread.utils.rowcol_to_a1(1, len(SHEET_COLUMNS))}"
    ws.format(header, {"textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                       "backgroundColor": {"red": 0.12, "green": 0.23, "blue": 0.37},
                       "horizontalAlignment": "CENTER"})
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}"
